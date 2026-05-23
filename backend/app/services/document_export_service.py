"""
AIuthor Backend — Document Export Service.

Handles assembly, DOCX generation, and PDF conversion of book manuscripts.
"""
from __future__ import annotations

import logging
import os
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path
from uuid import UUID, uuid4
from sqlalchemy.orm import Session

from app.config import get_settings
from app.workflows.schemas import (
    BookAssemblyRequest,
    BookAssemblyResponse,
    BookExportRequest,
    BookExportFileItem,
    BookExportResponse,
)
from app.schemas.export import ExportFileCreate
from app.services.book_assembler_service import BookAssemblerService
from app.services.export_service import ExportService
from app.services.exceptions import ValidationServiceError

# Setup python-docx import safely
try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    Document = None
    Pt = None
    WD_ALIGN_PARAGRAPH = None

logger = logging.getLogger(__name__)


class DocumentExportService:
    """
    Service to handle manuscript file generation, formatting, and DB persistence.
    """

    def __init__(self, db: Session) -> None:
        self.db = db
        self.assembler = BookAssemblerService(db)
        self.export_service = ExportService(db)
        self.settings = get_settings()

    def _get_output_dir(self, book_id: UUID, run_id: UUID | None = None) -> Path:
        """
        Determines the output directory for export files.
        Creates it if it does not exist.
        """
        base_dir = Path(self.settings.export_output_dir)
        sub_dir = run_id if run_id is not None else "manual"
        output_path = base_dir / str(book_id) / str(sub_dir)
        output_path.mkdir(parents=True, exist_ok=True)
        return output_path

    def _safe_filename(self, title: str, suffix: str) -> str:
        """
        Slugifies the book title and appends a timestamp to avoid name collisions.
        """
        # Slugify: lower, keep alpha-numeric and hyphens
        slug = title.strip().lower()
        slug = re.sub(r"\s+", "-", slug)
        slug = re.sub(r"[^\w\-]", "", slug)
        slug = re.sub(r"-+", "-", slug)
        if not slug:
            slug = "manuscript"
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"{slug}_{timestamp}.{suffix}"

    def generate_docx(self, assembly: BookAssemblyResponse, output_path: Path) -> Path:
        """
        Generates a structured DOCX manuscript from the assembly response.
        """
        if Document is None:
            raise ValidationServiceError(
                message="python-docx library is not installed.",
                code="python_docx_missing"
            )

        doc = Document()

        # 1. Title Page
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(assembly.title)
        title_run.font.name = "Arial"
        title_run.font.size = Pt(28)
        title_run.font.bold = True

        if assembly.subtitle:
            subtitle_p = doc.add_paragraph()
            subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle_run = subtitle_p.add_run(assembly.subtitle)
            subtitle_run.font.name = "Arial"
            subtitle_run.font.size = Pt(16)
            subtitle_run.font.italic = True

        # Spacer paragraphs
        for _ in range(5):
            doc.add_paragraph()

        author_p = doc.add_paragraph()
        author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author_run = author_p.add_run(f"By {assembly.author or 'AIuthor'}")
        author_run.font.name = "Arial"
        author_run.font.size = Pt(14)

        doc.add_page_break()

        # 2. Table of Contents Placeholder (Virtual front matter has it, but we render list of chapters here)
        if assembly.toc:
            doc.add_heading("Table of Contents", level=1)
            for item in assembly.toc:
                doc.add_paragraph(f"Chapter {item.get('chapter_number')}: {item.get('title')}")
            doc.add_page_break()

        # 3. Front Matter (Excluding title_page and toc as they are custom generated above)
        for section in assembly.front_matter:
            stype = section.get("section_type")
            if stype in ["title_page", "toc"]:
                continue
            title = section.get("title") or stype.replace("_", " ").title()
            doc.add_heading(title, level=1)
            content = section.get("content") or ""
            for p_text in content.split("\n"):
                if p_text.strip():
                    doc.add_paragraph(p_text.strip())
            doc.add_page_break()

        # 4. Chapters
        for chapter in assembly.chapters:
            doc.add_heading(f"Chapter {chapter.chapter_number}: {chapter.title}", level=1)
            content = chapter.content or ""
            for p_text in content.split("\n"):
                if p_text.strip():
                    doc.add_paragraph(p_text.strip())
            doc.add_page_break()

        # 5. Back Matter
        for section in assembly.back_matter:
            stype = section.get("section_type")
            title = section.get("title") or stype.replace("_", " ").title()
            doc.add_heading(title, level=1)
            content = section.get("content") or ""
            for p_text in content.split("\n"):
                if p_text.strip():
                    doc.add_paragraph(p_text.strip())
            doc.add_page_break()

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return output_path

    def convert_docx_to_pdf(self, docx_path: Path, pdf_path: Path) -> Path:
        """
        Converts DOCX to PDF using LibreOffice headless command line.
        """
        # Look for soffice executable
        soffice_path = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice_path:
            # Check standard Windows paths
            std_paths = [
                Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
                Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe")
            ]
            for p in std_paths:
                if p.exists():
                    soffice_path = str(p)
                    break

        if not soffice_path:
            raise ValidationServiceError(
                message="LibreOffice is not installed or not found in PATH. PDF export requires LibreOffice.",
                code="libreoffice_not_found"
            )

        outdir = pdf_path.parent
        cmd = [
            soffice_path,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(docx_path)
        ]
        try:
            subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=60, check=True)
            expected_pdf_name = docx_path.stem + ".pdf"
            generated_pdf_path = outdir / expected_pdf_name
            if generated_pdf_path.exists():
                if generated_pdf_path.resolve() != pdf_path.resolve():
                    if pdf_path.exists():
                        pdf_path.unlink()
                    generated_pdf_path.rename(pdf_path)
                return pdf_path
            else:
                raise ValidationServiceError(
                    message=f"PDF conversion completed but output file was not found: {generated_pdf_path}",
                    code="pdf_conversion_file_not_found"
                )
        except subprocess.SubprocessError as e:
            raise ValidationServiceError(
                message=f"LibreOffice PDF conversion failed: {str(e)}",
                code="pdf_conversion_failed"
            )

    def _create_export_file_record(
        self,
        book_id: UUID,
        run_id: UUID | None,
        export_type: str,
        status: str,
        file_path: Path | None,
        file_name: str | None,
        error_message: str | None = None,
        metadata: dict | None = None
    ) -> UUID:
        """
        Creates a row in export_files database table.
        """
        mime_types = {
            "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "pdf": "application/pdf"
        }
        
        file_size = None
        if file_path and file_path.exists():
            file_size = file_path.stat().st_size

        payload = ExportFileCreate(
            book_id=book_id,
            run_id=run_id,
            export_type=export_type,
            file_path=str(file_path) if file_path else f"failed://exports/{book_id}/{export_type}",
            file_name=file_name or f"{export_type}_{book_id}.failed",
            mime_type=mime_types.get(export_type, "application/octet-stream"),
            status=status,
            export_metadata={
                **(metadata or {}),
                "error_message": error_message,
                "file_size_bytes": file_size
            }
        )
        
        db_file = self.export_service.create_export_file(
            payload=payload,
            book_id=book_id,
            run_id=run_id
        )
        return db_file.id

    def export_book(self, request: BookExportRequest) -> BookExportResponse:
        """
        Assembles a book and generates requested export files synchronously.
        """
        # 1. Assemble manuscript first
        assembly_req = BookAssemblyRequest(
            book_id=request.book_id,
            run_id=request.run_id,
            include_front_matter=request.include_front_matter,
            include_back_matter=request.include_back_matter,
            include_toc=request.include_toc,
            include_glossary=request.include_glossary,
            include_bibliography=request.include_bibliography,
            include_memory_notes=request.include_memory_notes,
            prefer_final_text=getattr(request, "prefer_final_text", True),
            metadata=request.metadata,
        )

        assembly = self.assembler.assemble_book(assembly_req)

        output_dir = self._get_output_dir(request.book_id, request.run_id)
        
        results = []
        docx_generated_path = None
        docx_file_name = None
        
        # Determine files to generate
        export_types = request.export_types
        
        # We always need DOCX generated first if PDF is requested
        generate_docx_format = "docx" in export_types or "pdf" in export_types
        
        if generate_docx_format:
            docx_file_name = self._safe_filename(assembly.title, "docx")
            docx_generated_path = output_dir / docx_file_name
            
            try:
                self.generate_docx(assembly, docx_generated_path)
                
                # If docx was explicitly requested, save DB entry now
                if "docx" in export_types:
                    export_id = self._create_export_file_record(
                        book_id=request.book_id,
                        run_id=request.run_id,
                        export_type="docx",
                        status="ready",
                        file_path=docx_generated_path,
                        file_name=docx_file_name,
                        metadata=request.metadata
                    )
                    results.append(BookExportFileItem(
                        export_id=export_id,
                        export_type="docx",
                        status="ready",
                        file_name=docx_file_name,
                        file_path=str(docx_generated_path),
                        file_size_bytes=docx_generated_path.stat().st_size
                    ))
            except Exception as e:
                logger.error("DOCX generation failed: %s", str(e), exc_info=True)
                if "docx" in export_types:
                    export_id = self._create_export_file_record(
                        book_id=request.book_id,
                        run_id=request.run_id,
                        export_type="docx",
                        status="failed",
                        file_path=None,
                        file_name=docx_file_name,
                        error_message=str(e),
                        metadata=request.metadata
                    )
                    results.append(BookExportFileItem(
                        export_id=export_id,
                        export_type="docx",
                        status="failed",
                        file_name=docx_file_name,
                        error_message=str(e)
                    ))
        
        if "pdf" in export_types:
            pdf_file_name = self._safe_filename(assembly.title, "pdf")
            pdf_generated_path = output_dir / pdf_file_name
            
            if not self.settings.enable_pdf_export:
                export_id = self._create_export_file_record(
                    book_id=request.book_id,
                    run_id=request.run_id,
                    export_type="pdf",
                    status="failed",
                    file_path=None,
                    file_name=pdf_file_name,
                    error_message="PDF export is disabled in configuration.",
                    metadata=request.metadata
                )
                results.append(BookExportFileItem(
                    export_id=export_id,
                    export_type="pdf",
                    status="failed",
                    file_name=pdf_file_name,
                    error_message="PDF export is disabled in configuration."
                ))
            elif not docx_generated_path or not docx_generated_path.exists():
                err_msg = "Cannot generate PDF because DOCX generation failed or was skipped."
                export_id = self._create_export_file_record(
                    book_id=request.book_id,
                    run_id=request.run_id,
                    export_type="pdf",
                    status="failed",
                    file_path=None,
                    file_name=pdf_file_name,
                    error_message=err_msg,
                    metadata=request.metadata
                )
                results.append(BookExportFileItem(
                    export_id=export_id,
                    export_type="pdf",
                    status="failed",
                    file_name=pdf_file_name,
                    error_message=err_msg
                ))
            else:
                try:
                    self.convert_docx_to_pdf(docx_generated_path, pdf_generated_path)
                    export_id = self._create_export_file_record(
                        book_id=request.book_id,
                        run_id=request.run_id,
                        export_type="pdf",
                        status="ready",
                        file_path=pdf_generated_path,
                        file_name=pdf_file_name,
                        metadata=request.metadata
                    )
                    results.append(BookExportFileItem(
                        export_id=export_id,
                        export_type="pdf",
                        status="ready",
                        file_name=pdf_file_name,
                        file_path=str(pdf_generated_path),
                        file_size_bytes=pdf_generated_path.stat().st_size
                    ))
                except Exception as e:
                    logger.error("PDF conversion failed: %s", str(e), exc_info=True)
                    export_id = self._create_export_file_record(
                        book_id=request.book_id,
                        run_id=request.run_id,
                        export_type="pdf",
                        status="failed",
                        file_path=None,
                        file_name=pdf_file_name,
                        error_message=str(e),
                        metadata=request.metadata
                    )
                    results.append(BookExportFileItem(
                        export_id=export_id,
                        export_type="pdf",
                        status="failed",
                        file_name=pdf_file_name,
                        error_message=str(e)
                    ))
                    
        # Compute overall status
        statuses = [item.status for item in results]
        if not statuses:
            overall_status = "failed"
        elif all(s == "ready" for s in statuses):
            overall_status = "completed"
        elif all(s == "failed" for s in statuses):
            overall_status = "failed"
        else:
            overall_status = "partial_failed"

        return BookExportResponse(
            book_id=request.book_id,
            run_id=request.run_id,
            status=overall_status,
            assembly=assembly,
            files=results,
            metadata=request.metadata
        )
